#!/usr/bin/env python3
"""Extract CPF numbers from DOCX files in the current directory.

The CPF (Cadastro de Pessoas Físicas) can appear in two formats:
    - ``XXX.XXX.XXX-YY`` (with punctuation)
    - ``XXXXXXXXXXX`` (11 continuous digits)

This script scans every ``.docx`` file in the repository directory and prints the
file name followed by any CPF numbers found within that file.
"""

from __future__ import annotations

import glob
import re
from pathlib import Path
from typing import Iterable, List

from docx import Document

CPF_PATTERN = re.compile(r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b|\b\d{11}\b")


def _extract_text_from_doc(path: Path) -> str:
    """Return the full text content of a DOCX file.

    Text is collected from paragraphs and table cells to maximise coverage of the
    document contents.
    """

    doc = Document(str(path))
    parts: List[str] = []

    # Paragraph text
    for paragraph in doc.paragraphs:
        parts.append(paragraph.text)

    # Table text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)

    return "\n".join(parts)


def extract_cpfs(path: Path) -> List[str]:
    """Extract CPF numbers from the DOCX file at *path*.

    Returns a list of all CPF strings found. Duplicate matches within a file are
    not deduplicated.
    """

    text = _extract_text_from_doc(path)
    return CPF_PATTERN.findall(text)


def main() -> None:
    docx_files: Iterable[Path] = sorted(Path('.').glob('*.docx'))
    for doc_path in docx_files:
        cpfs = extract_cpfs(doc_path)
        if cpfs:
            print(f"{doc_path.name}: {', '.join(cpfs)}")
        else:
            print(f"{doc_path.name}: CPF not found")


if __name__ == "__main__":
    main()
